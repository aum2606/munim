import io

import pytest
from docx import Document
from pypdf import PdfWriter

from app.core.errors import DomainError
from app.db.models.knowledge import SourceType
from app.ingestion.chunker import chunk_pages
from app.ingestion.parsers import PageText, parse_document


def test_overlapping_chunks_respect_headings_and_pages() -> None:
    chunks = chunk_pages(
        [
            PageText(
                "# First\n\n" + " ".join(str(n) for n in range(12)) + "\n\n# Second\n\nNew section",
                1,
            ),
            PageText("Other page", 2),
        ],
        size=5,
        overlap=1,
    )
    assert [chunk.content for chunk in chunks[:3]] == ["0 1 2 3 4", "4 5 6 7 8", "8 9 10 11"]
    assert chunks[3].section == "Second"
    assert chunks[4].page == 2
    assert chunk_pages([PageText(" ", 1)]) == []
    with pytest.raises(ValueError):
        chunk_pages([], size=3, overlap=3)


def test_txt_docx_and_empty_pdf_parsing() -> None:
    assert parse_document(b"Hello", SourceType.TXT)[0].content == "Hello"
    with pytest.raises(DomainError):
        parse_document(b"\xff", SourceType.TXT)
    document = Document()
    document.add_heading("Warranty", 1)
    document.add_paragraph("One year warranty")
    buffer = io.BytesIO()
    document.save(buffer)
    text = parse_document(buffer.getvalue(), SourceType.DOCX)[0].content
    assert "# Warranty" in text and "One year warranty" in text
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    assert parse_document(buffer.getvalue(), SourceType.PDF)[0].content == ""


async def test_hosted_embedder_uses_mock_transport_and_validates_dimension() -> None:
    import httpx

    from app.core.config import Settings
    from app.llm.embedder import OpenAICompatibleEmbedder

    settings = Settings(
        secret_key="test-only-key-with-at-least-32-characters",
        embedding_model="synthetic-embedding",
        embedding_dim=2,
        llm_api_key="fake-key",
    )

    async def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path == "/v1/embeddings"
        return httpx.Response(200, json={"data": [{"index": 0, "embedding": [1, 0]}]})

    embedder = OpenAICompatibleEmbedder(settings, httpx.MockTransport(handler))
    assert len((await embedder.embed(["hello"]))[0]) == 2
    assert await embedder.embed([]) == []
